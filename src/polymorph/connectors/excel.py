from __future__ import annotations

import math
import re
import zipfile
from collections import Counter
from collections.abc import Iterable, Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import BinaryIO, Protocol, cast

import openpyxl
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from polymorph.classification import classify_field_name, infer_role
from polymorph.content import (
    ContentInspector,
    ContentKind,
    FileIdentity,
    require_matching_file_identity,
)
from polymorph.errors import ConnectorError
from polymorph.matching.deterministic import normalize_name
from polymorph.models.schema import FieldDescriptor, SchemaDescriptor
from polymorph.models.types import DataType, Sensitivity

from .base import ConnectorCapabilities

_FIXED_ZERO_FORMAT = re.compile(r"^0+$")
_FORMULA_TAG = re.compile(rb"<f(?:\s|>)")


class _Cell(Protocol):
    @property
    def value(self) -> object: ...

    @property
    def number_format(self) -> str: ...


class _Worksheet(Protocol):
    title: str
    max_row: int | None
    max_column: int | None

    def iter_rows(
        self,
        *,
        min_row: int | None = None,
        max_row: int | None = None,
        min_col: int | None = None,
        max_col: int | None = None,
    ) -> Iterator[tuple[_Cell, ...]]: ...


@dataclass(frozen=True, slots=True)
class ExcelLayout:
    header_row: int
    data_start_row: int
    width: int
    confidence: float
    reasons: tuple[str, ...] = ()


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _row_values(row: Iterable[_Cell]) -> list[object]:
    return [cell.value for cell in row]


def _trim_width(values: list[object]) -> int:
    width = 0
    for index, value in enumerate(values, start=1):
        if not _is_empty(value):
            width = index
    return width


def _header_score(
    values: list[object], following: list[list[object]]
) -> tuple[float, tuple[str, ...]]:
    width = _trim_width(values)
    if width == 0:
        return 0.0, ()
    active = values[:width]
    nonempty = [value for value in active if not _is_empty(value)]
    if not nonempty:
        return 0.0, ()

    text = [value for value in nonempty if isinstance(value, str) and value.strip()]
    text_ratio = len(text) / len(nonempty)
    normalized = [normalize_name(str(value)) for value in nonempty]
    unique_ratio = len(set(normalized)) / len(normalized) if normalized else 0.0
    coverage = len(nonempty) / width

    next_densities: list[float] = []
    next_nontext_ratios: list[float] = []
    for row in following[:3]:
        sample = row[:width]
        present = [value for value in sample if not _is_empty(value)]
        if not present:
            continue
        next_densities.append(len(present) / width)
        next_nontext_ratios.append(
            sum(not isinstance(value, str) for value in present) / len(present)
        )
    next_density = sum(next_densities) / len(next_densities) if next_densities else 0.0
    data_transition = (
        sum(next_nontext_ratios) / len(next_nontext_ratios) if next_nontext_ratios else 0.0
    )

    # A one-cell title should not beat a wider structured row, but single-column files are
    # still supported when the following rows have comparable density.
    structure_bonus = min(len(nonempty), 8) / 8
    score = (
        0.30 * text_ratio
        + 0.18 * unique_ratio
        + 0.17 * coverage
        + 0.23 * next_density
        + 0.07 * data_transition
        + 0.05 * structure_bonus
    )
    if len(nonempty) == 1 and width == 1:
        score *= 0.92
    elif len(nonempty) == 1:
        score *= 0.45

    reasons = []
    if text_ratio >= 0.8:
        reasons.append("mostly textual labels")
    if unique_ratio >= 0.9:
        reasons.append("labels are unique")
    if next_density >= 0.6:
        reasons.append("dense rows follow candidate")
    if data_transition >= 0.35:
        reasons.append("following rows look like data")
    return score, tuple(reasons)


def _type_of(value: object) -> DataType:
    if isinstance(value, bool):
        return DataType.BOOLEAN
    if isinstance(value, datetime):
        return DataType.DATETIME
    if isinstance(value, date):
        return DataType.DATE
    if isinstance(value, int):
        return DataType.INTEGER
    if isinstance(value, (float, Decimal)):
        return DataType.DECIMAL
    if isinstance(value, bytes):
        return DataType.BINARY
    if isinstance(value, str):
        return DataType.STRING
    return DataType.UNKNOWN


def _merge_types(types: list[DataType]) -> DataType:
    known = [item for item in types if item is not DataType.UNKNOWN]
    if not known:
        return DataType.UNKNOWN
    distinct = set(known)
    if len(distinct) == 1:
        return known[0]
    if distinct <= {DataType.INTEGER, DataType.DECIMAL}:
        return DataType.DECIMAL
    if distinct <= {DataType.DATE, DataType.DATETIME}:
        return DataType.DATETIME
    # Mixed text and typed values are safest as strings at the schema boundary. The
    # mapping plan may still choose an explicit parser after review.
    if DataType.STRING in distinct:
        return DataType.STRING
    return DataType.UNKNOWN


def _value_from_cell(cell: _Cell) -> object:
    value = cell.value
    if isinstance(value, int) and not isinstance(value, bool):
        number_format = str(getattr(cell, "number_format", "") or "")
        if _FIXED_ZERO_FORMAT.fullmatch(number_format):
            return f"{value:0{len(number_format)}d}"
    if isinstance(value, float) and not math.isfinite(value):
        return value
    return value


class ExcelConnector:
    capabilities = ConnectorCapabilities(read_schema=True, read_records=True)

    def __init__(
        self,
        path: str | Path,
        sheet: str | None = None,
        header_row: int | None = None,
        *,
        sensitivity_overrides: Mapping[str, Sensitivity] | None = None,
        layout_scan_rows: int = 64,
        type_sample_rows: int = 128,
        max_columns: int = 512,
        skip_repeated_headers: bool = True,
        expected_source_identity: FileIdentity | None = None,
    ) -> None:
        self.path = Path(path)
        self.sheet = sheet
        self.header_row = header_row
        self.sensitivity_overrides = dict(sensitivity_overrides or {})
        self.layout_scan_rows = max(4, layout_scan_rows)
        self.type_sample_rows = max(1, type_sample_rows)
        self.max_columns = max(1, max_columns)
        self.skip_repeated_headers = skip_repeated_headers
        self._layout_cache: ExcelLayout | None = None
        self._schema_cache: SchemaDescriptor | None = None
        self._content_inspector = ContentInspector()
        self._trusted_identity: FileIdentity | None = None
        self._expected_source_identity = expected_source_identity

    def _ensure_source_safe(self) -> FileIdentity:
        if self._trusted_identity is None:
            report = self._content_inspector.require(
                self.path,
                allowed={ContentKind.XLSX},
                purpose="Excel connector",
            )
            if (
                self._expected_source_identity is not None
                and report.identity != self._expected_source_identity
            ):
                raise ConnectorError(
                    "Excel connector rejected input changed after accepted content inspection"
                )
            self._trusted_identity = report.identity
        return self._trusted_identity

    def _open_source_handle(self) -> BinaryIO:
        expected = self._ensure_source_safe()
        try:
            source_handle = self.path.open("rb")
        except OSError as exc:
            raise ConnectorError("Excel connector could not open the inspected input") from exc
        try:
            require_matching_file_identity(source_handle, expected, purpose="Excel connector")
        except Exception:
            source_handle.close()
            raise
        return source_handle

    def _close_source_handle(self, source_handle: BinaryIO) -> None:
        try:
            require_matching_file_identity(
                source_handle,
                self._ensure_source_safe(),
                purpose="Excel connector",
            )
        finally:
            source_handle.close()

    def _verify_cached_source(self) -> None:
        self._close_source_handle(self._open_source_handle())

    def _worksheet(self, *, data_only: bool) -> tuple[Workbook, _Worksheet, BinaryIO]:
        if not bool(getattr(openpyxl, "DEFUSEDXML", False)):
            raise ConnectorError(
                "Excel connector requires openpyxl XML hardening through defusedxml"
            )
        source_handle = self._open_source_handle()
        try:
            workbook = load_workbook(
                source_handle,
                read_only=True,
                data_only=data_only,
                keep_links=False,
                keep_vba=False,
            )
        except Exception:
            source_handle.close()
            raise
        worksheet = workbook[self.sheet] if self.sheet else workbook.active
        if worksheet is None or not hasattr(worksheet, "iter_rows"):
            workbook.close()
            source_handle.close()
            raise ConnectorError("selected Excel sheet is not a worksheet")
        return workbook, cast(_Worksheet, worksheet), source_handle

    def _close_workbook(self, workbook: Workbook, source_handle: BinaryIO) -> None:
        try:
            require_matching_file_identity(
                source_handle,
                self._ensure_source_safe(),
                purpose="Excel connector",
            )
        finally:
            try:
                workbook.close()
            finally:
                source_handle.close()

    def _has_formulas(self) -> bool:
        """Detect formulas in the selected worksheet without trusting cached values.

        openpyxl deliberately does not calculate formulas. Reading with ``data_only=True``
        therefore returns workbook-provided cached results, which may be stale. We scan the
        selected OOXML worksheet stream once and surface that uncertainty to preflight.
        """

        workbook, worksheet, source_handle = self._worksheet(data_only=False)
        try:
            member = getattr(worksheet, "_worksheet_path", None)
        finally:
            self._close_workbook(workbook, source_handle)
        if not isinstance(member, str) or not member:
            return True  # unknown provenance is safer than silently assuming no formulas

        archive_handle = self._open_source_handle()
        try:
            with zipfile.ZipFile(archive_handle) as archive, archive.open(member, "r") as stream:
                overlap = b""
                while True:
                    chunk = stream.read(64 * 1024)
                    if not chunk:
                        return False
                    probe = overlap + chunk
                    if _FORMULA_TAG.search(probe):
                        return True
                    overlap = probe[-8:]
        except (OSError, KeyError, RuntimeError, zipfile.BadZipFile):
            return True
        finally:
            self._close_source_handle(archive_handle)

    def discover_layout(self) -> ExcelLayout:
        if self._layout_cache is not None:
            self._verify_cached_source()
            return self._layout_cache

        workbook, worksheet, source_handle = self._worksheet(data_only=False)
        try:
            max_row = worksheet.max_row or 1
            max_column = worksheet.max_column or 1
            rows = [
                _row_values(row)
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=min(self.layout_scan_rows, max_row),
                    max_col=min(self.max_columns, max_column),
                )
            ]
            if not rows:
                raise ValueError("Excel worksheet is empty")

            if self.header_row is not None:
                if self.header_row < 1 or self.header_row > len(rows):
                    raise ValueError("configured Excel header row is outside the scanned sheet")
                width = _trim_width(rows[self.header_row - 1])
                if width == 0:
                    raise ValueError("configured Excel header row is empty")
                layout = ExcelLayout(
                    header_row=self.header_row,
                    data_start_row=self.header_row + 1,
                    width=width,
                    confidence=1.0,
                    reasons=("explicit header row",),
                )
                self._layout_cache = layout
                return layout

            candidates: list[tuple[float, int, int, tuple[str, ...]]] = []
            for index, values in enumerate(rows):
                score, reasons = _header_score(values, rows[index + 1 : index + 4])
                width = _trim_width(values)
                if width:
                    candidates.append((score, index + 1, width, reasons))
            if not candidates:
                raise ValueError("no structured header row could be detected")

            candidates.sort(key=lambda item: (item[0], item[2], -item[1]), reverse=True)
            best_score, _, best_width, _ = candidates[0]
            # Paged exports frequently repeat the same header later in the sheet. When
            # several structured rows are essentially tied, the earliest one defines the
            # dataset boundary and later copies are handled as repeated headers.
            competitive = [
                item
                for item in candidates
                if item[0] >= best_score - 0.03 and item[2] >= max(1, int(best_width * 0.8))
            ]
            score, row_number, width, reasons = min(competitive, key=lambda item: item[1])
            if score < 0.42:
                raise ValueError("Excel layout is too ambiguous for automatic header detection")
            runner_up = max((item[0] for item in candidates if item[1] != row_number), default=0.0)
            confidence = max(0.0, min(1.0, score * 0.8 + max(0.0, score - runner_up) * 0.7))
            layout = ExcelLayout(
                header_row=row_number,
                data_start_row=row_number + 1,
                width=width,
                confidence=confidence,
                reasons=reasons,
            )
            self._layout_cache = layout
            return layout
        finally:
            self._close_workbook(workbook, source_handle)

    def _headers(self, worksheet: _Worksheet, layout: ExcelLayout) -> list[str]:
        row = next(
            worksheet.iter_rows(
                min_row=layout.header_row,
                max_row=layout.header_row,
                max_col=layout.width,
            )
        )
        headers: list[str] = []
        seen: Counter[str] = Counter()
        for index, cell in enumerate(row, start=1):
            raw = cell.value
            name = str(raw).strip() if not _is_empty(raw) else f"column_{index}"
            seen[name] += 1
            if seen[name] > 1:
                name = f"{name} [{seen[name]}]"
            headers.append(name)
        return headers

    def inspect_schema(self) -> SchemaDescriptor:
        if self._schema_cache is not None:
            self._verify_cached_source()
            return self._schema_cache
        layout = self.discover_layout()
        workbook, worksheet, source_handle = self._worksheet(data_only=True)
        try:
            headers = self._headers(worksheet, layout)
            samples: list[list[object]] = [[] for _ in range(layout.width)]
            nulls = [False] * layout.width
            sampled = 0
            max_row = worksheet.max_row or layout.data_start_row
            for row in worksheet.iter_rows(
                min_row=layout.data_start_row,
                max_row=min(
                    max_row,
                    layout.data_start_row + self.type_sample_rows - 1,
                ),
                max_col=layout.width,
            ):
                values = [_value_from_cell(cell) for cell in row]
                if not any(not _is_empty(value) for value in values):
                    continue
                if self._is_repeated_header(values, headers):
                    continue
                sampled += 1
                for index, value in enumerate(values):
                    if _is_empty(value):
                        nulls[index] = True
                    else:
                        samples[index].append(value)

            fields = []
            for index, name in enumerate(headers, start=1):
                observed_types = [_type_of(value) for value in samples[index - 1]]
                fields.append(
                    FieldDescriptor(
                        id=f"c{index}",
                        name=name,
                        data_type=_merge_types(observed_types),
                        nullable=nulls[index - 1] or sampled == 0,
                        sensitivity=self.sensitivity_overrides.get(name, classify_field_name(name)),
                        role=infer_role(name),
                        container=worksheet.title,
                    )
                )
            formulas_present = self._has_formulas()
            schema = SchemaDescriptor(
                id=f"excel:{self.path.name}:{worksheet.title}",
                fields=tuple(fields),
                metadata={
                    "path": self.path.name,
                    "sheet": worksheet.title,
                    "header_row": str(layout.header_row),
                    "layout_confidence": f"{layout.confidence:.4f}",
                    "formula_cells_present": "true" if formulas_present else "false",
                    "formula_value_source": "cached_workbook_value"
                    if formulas_present
                    else "literal",
                },
            )
            self._schema_cache = schema
            return schema
        finally:
            self._close_workbook(workbook, source_handle)

    @staticmethod
    def _is_repeated_header(values: list[object], headers: list[str]) -> bool:
        comparable = 0
        matches = 0
        for value, header in zip(values, headers, strict=False):
            if _is_empty(value):
                continue
            comparable += 1
            if normalize_name(str(value)) == normalize_name(header.split(" [", 1)[0]):
                matches += 1
        return comparable >= 2 and matches / comparable >= 0.8

    def iter_records(self) -> Iterable[Mapping[str, object]]:
        layout = self.discover_layout()
        schema = self.inspect_schema()
        headers = [field.name for field in schema.fields]
        workbook, worksheet, source_handle = self._worksheet(data_only=True)
        try:
            for row in worksheet.iter_rows(
                min_row=layout.data_start_row,
                max_col=layout.width,
            ):
                values = [_value_from_cell(cell) for cell in row]
                if not any(not _is_empty(value) for value in values):
                    continue
                if self.skip_repeated_headers and self._is_repeated_header(values, headers):
                    continue
                yield {f"c{index}": value for index, value in enumerate(values, start=1)}
        finally:
            self._close_workbook(workbook, source_handle)
