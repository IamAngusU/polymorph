import re
import zipfile

import pytest
from openpyxl import Workbook

from polymorph.connectors.excel import ExcelConnector
from polymorph.errors import ConnectorError
from polymorph.models.types import DataType, FieldRole, Sensitivity


def _remove_worksheet_dimension(path) -> None:
    rewritten = path.with_name("rewritten.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload, replacements = re.subn(rb"<dimension[^>]*/>", b"", payload, count=1)
                assert replacements == 1
            target.writestr(info, payload)
    rewritten.replace(path)


def test_excel_schema_and_records(tmp_path):
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Debitor Nr", "API Token", "Netto Betrag"])
    sheet.append(["00042", "secret", 12.5])
    workbook.save(path)

    connector = ExcelConnector(path)
    schema = connector.inspect_schema()
    records = list(connector.iter_records())

    assert schema.fields[1].sensitivity is Sensitivity.SECRET
    assert schema.fields[0].role is FieldRole.NATURAL_KEY
    assert schema.fields[2].data_type is DataType.DECIMAL
    assert records[0]["c1"] == "00042"
    assert records[0]["c3"] == 12.5


def test_excel_auto_detects_header_below_title_rows(tmp_path):
    path = tmp_path / "messy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Quarterly customer export"
    sheet.append([])
    sheet.append(["Customer Number", "Order Number", "Net Amount"])
    sheet.append([42, "A-1", 100.25])
    sheet.append([43, "A-2", 99.50])
    workbook.save(path)

    connector = ExcelConnector(path)
    layout = connector.discover_layout()
    schema = connector.inspect_schema()

    assert layout.header_row == 3
    assert [field.name for field in schema.fields] == [
        "Customer Number",
        "Order Number",
        "Net Amount",
    ]
    assert schema.fields[0].data_type is DataType.INTEGER
    assert schema.fields[2].data_type is DataType.DECIMAL


def test_excel_recovers_missing_dimensions_before_header_detection(tmp_path):
    path = tmp_path / "unsized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Synthetic export"])
    sheet.append([])
    sheet.append([])
    sheet.append(["Order", "Email", "Quantity", "Amount", "Token", "Note"])
    sheet.append(["A-1", "a@example.invalid", 2, 199, "token-1", "first"])
    sheet.append(["A-2", "b@example.invalid", 3, 299, "token-2", "second"])
    workbook.save(path)
    _remove_worksheet_dimension(path)

    connector = ExcelConnector(path)
    layout = connector.discover_layout()
    records = list(connector.iter_records())

    assert layout.header_row == 4
    assert layout.width == 6
    assert len(records) == 2
    assert records[0]["c6"] == "first"


def test_excel_preserves_fixed_width_numeric_identifiers(tmp_path):
    path = tmp_path / "ids.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Customer Number", "Name"])
    sheet.append([42, "Acme"])
    sheet["A2"].number_format = "000000"
    workbook.save(path)

    connector = ExcelConnector(path)
    record = next(iter(connector.iter_records()))
    assert record["c1"] == "000042"
    assert connector.inspect_schema().fields[0].data_type is DataType.STRING


def test_excel_skips_repeated_headers_inside_export(tmp_path):
    path = tmp_path / "paged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Customer Number", "Amount"])
    sheet.append(["A-1", 10])
    sheet.append(["Customer Number", "Amount"])
    sheet.append(["A-2", 20])
    workbook.save(path)

    records = list(ExcelConnector(path).iter_records())
    assert len(records) == 2
    assert [record["c1"] for record in records] == ["A-1", "A-2"]


def test_excel_marks_formula_values_as_cached_and_unproven(tmp_path):
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Amount", "Tax"])
    sheet.append([100, "=A2*0.19"])
    workbook.save(path)

    schema = ExcelConnector(path).inspect_schema()

    assert schema.metadata["formula_cells_present"] == "true"
    assert schema.metadata["formula_value_source"] == "cached_workbook_value"


def test_excel_fails_closed_when_openpyxl_xml_hardening_is_disabled(tmp_path, monkeypatch) -> None:
    import polymorph.connectors.excel as excel_module

    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active.append(["id"])
    workbook.active.append([1])
    workbook.save(path)

    monkeypatch.setattr(excel_module.openpyxl, "DEFUSEDXML", False)

    with pytest.raises(ConnectorError, match="defusedxml"):
        ExcelConnector(path).inspect_schema()


def test_excel_rejects_path_swap_after_content_gate(tmp_path) -> None:
    path = tmp_path / "source.xlsx"
    replacement = tmp_path / "replacement.xlsx"
    original_workbook = Workbook()
    original_workbook.active.append(["name", "value"])
    original_workbook.active.append(["Alice", 1])
    original_workbook.save(path)
    replacement_workbook = Workbook()
    replacement_workbook.active.append(["name", "value"])
    replacement_workbook.active.append(["Mallory", 999])
    replacement_workbook.save(replacement)
    connector = ExcelConnector(path)
    connector.inspect_schema()

    replacement.replace(path)

    with pytest.raises(ConnectorError, match="changed after content inspection"):
        connector.inspect_schema()
    with pytest.raises(ConnectorError, match="changed after content inspection"):
        list(connector.iter_records())
